#!/usr/bin/env python3
"""Generate the thesis proposal and literature review from local paper assets."""

from __future__ import annotations

import csv
import shutil
import sys
from pathlib import Path

DEPS = Path("/tmp/topic_pydeps")
if DEPS.exists():
    sys.path.insert(0, str(DEPS))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper_assets" / "topic"
TEMPLATE = OUT / "template"
PROPOSAL_TEMPLATE = TEMPLATE / "1710105829_202642716378_研究生学位（毕业）论文开题报告2026.4.13.docx"
REVIEW_TEMPLATE = TEMPLATE / "1710105829_2026427163724_研究生学位（毕业）论文文献综述2026.4.13.docx"

ZH_TITLE = "异构联邦微调中面向资源约束的矩阵级LoRA调度研究"
EN_TITLE = "Resource-Constrained Matrix-Level LoRA Scheduling for Heterogeneous Federated Fine-Tuning"


PROPOSAL_ABSTRACT = (
    "随着预训练基础模型规模持续增长，终端设备和行业机构在不汇集原始数据的条件下协同微调模型，面临数据分布、训练显存、计算能力和通信带宽同时异构的难题。现有联邦LoRA方法主要从秩、物理层覆盖或量化精度调节训练容量，通常未将同一层内不同目标矩阵的效用差异作为独立调度维度；同时，统一的服务端重要性排序难以反映非独立同分布数据下的客户端个性化需求。本文拟研究面向异构客户端的矩阵级LoRA联邦微调方法FedMATRIX-LoRA。首先，客户端利用当前全局适配器和一个本地小批次计算矩阵级FIM重要性，实现不上传原始数据的客户端条件化画像；其次，建立由固定占用、LoRA参数与优化器状态以及位置相关激活组成的训练显存模型，并通过客户端动态规划生成满足显存和时延约束的紧凑候选配置；最后，服务端在轮次预算和时延目标下求解客户端—配置联合调度问题，选择每个客户端训练的固定秩LoRA矩阵，并采用通用联邦聚合完成更新。研究将在CIFAR100、GLUE四项任务和SAMSum上，从任务效果、累计矩阵预算、估计训练时间和异构适应性等方面，与FedAvg-LoRA、Uniform、Fed-HeLLo、LEGEND、HETLoRA及FAH-QLoRA等方法进行比较。现有实验已形成覆盖图像分类、语言理解和文本摘要的原型与结果台账，初步表明矩阵位置是独立且有效的资源分配维度。本研究旨在提高受限资源下联邦参数高效微调的配置灵活性和资源利用效率。"
)

SIGNIFICANCE = """1、研究意义（工程价值）

基础模型正逐步进入移动终端、边缘计算和跨机构智能应用。医疗、金融、政务、通信和个人终端中的训练数据往往具有隐私敏感、分散保存和难以集中等特点，联邦学习因而成为基础模型适配的重要技术路径。然而，基础模型即使采用参数高效微调，客户端仍需保存冻结主干、执行前向与反向传播并保留中间激活，且不同设备的训练显存、算力和带宽差异明显。使用统一LoRA配置容易造成低资源客户端不可训练、高资源客户端能力闲置或同步轮次被慢客户端拖延。

LoRA将每个目标线性变换的更新表示为低秩矩阵乘积，显著减少可训练参数和通信量，但“参数少”并不等价于“配置合理”。Transformer同一物理层中的Q、K、V、输出投影及前馈矩阵具有不同功能；在相同矩阵数量和秩下，不同位置组合可能产生不同任务收益。将物理层作为不可分割单元会迫使系统同时训练高效用与低效用矩阵，降低有限预算的利用率。因此，本研究将LoRA目标矩阵身份提升为资源调度变量，使系统能够在相同训练容量下选择更有价值的矩阵组合。

本课题具有三方面工程价值。第一，在原始数据不离开客户端的前提下，用本地数据估计矩阵效用，适配非独立同分布数据。第二，将位置相关训练显存、计算时延和通信时延纳入配置可行性判断，降低异构设备训练失败和同步等待风险。第三，通过客户端候选压缩与服务端全局调度，将指数规模的联合选择转化为可实现的两级优化流程，为边缘侧基础模型微调提供可复现的系统原型。"""

STATE_OF_ART = """2、国内外研究现状（工程应用现状）

（1）联邦学习与异构资源调度。FedAvg奠定了分布式本地训练与服务端聚合的基本范式。此后，Oort从统计效用与系统速度出发选择参与者，FedBalancer联合控制样本与训练节奏，PyramidFL、AutoFL和Cost-Effective FL进一步研究客户端选择、能耗与移动边缘资源约束。这类工作说明系统异构直接影响训练效率，但通常假设每个参与客户端采用预先确定的模型或训练配置，未细化到LoRA内部矩阵。

（2）参数高效微调。Adapter、Prefix-Tuning、LoRA和QLoRA分别通过插入瓶颈模块、优化连续提示、低秩更新和量化冻结主干降低微调成本。LoRA因结构简单、可与Transformer线性层自然结合而成为联邦基础模型微调的重要载体。现有研究还关注LoRA初始化、聚合偏差、精确聚合和多适配器对齐，例如FedEx-LoRA、LoRA-FAIR和TLoRA等。它们主要改进参数表示与聚合，而本课题关注聚合前“每个客户端应训练哪些目标矩阵”。

（3）异构联邦LoRA。FedPETuning、FedPEFT和相关研究将PEFT引入联邦环境。HeteroFL和FedRolex通过子模型宽度或滚动参数覆盖支持模型异构；FlexLoRA、HetLoRA和LEGEND主要调节客户端秩或训练容量；FAH-QLoRA联合量化精度与秩；Fed-HeLLo根据服务端代理数据上的层重要性，为不同显存客户端分配LoRA物理层。这些工作证明结构异构可以缓解资源差异，但仍存在两点空间：一是层级决策无法在同层内保留高价值投影并排除低价值投影；二是基于服务端代理数据的统一排序依赖代理数据的可获得性和代表性，在non-IID场景下难以表达客户端特定效用。

（4）训练显存与调度建模。训练显存除模型参数外，还包括梯度、优化器状态和中间激活。对部分层或部分矩阵反向传播时，最早被训练位置会影响需要保留的激活路径，因此同样数量的矩阵不一定具有相同训练显存代价。现有联邦调度往往将资源成本近似为参数量、层数或总计算量，缺少面向矩阵位置的可行配置生成及其与全局预算调度的结合。

综上，国内外研究已经形成联邦学习、PEFT、异构LoRA和资源调度四条技术线，但客户端本地数据驱动的矩阵级效用估计、位置相关训练显存约束和全局客户端—配置联合调度尚未被统一解决。本课题拟在不改变LoRA基本结构和联邦聚合范式的前提下，补足这一细粒度资源分配环节。"""

RESEARCH_PLAN = """本研究拟构建一个面向异构联邦微调的矩阵级LoRA资源调度框架。在不上传客户端原始数据的条件下，系统根据局部数据所反映的适配需求以及设备资源能力，选择每轮真正值得训练且能够训练的LoRA目标矩阵，使有限预算获得更高任务效用，同时兼顾训练显存可行性与同步轮次时延。

围绕这一目标，首先需要把可注入LoRA的目标线性矩阵从模型结构中动态识别出来，并将客户端每轮训练的矩阵子集表示为明确的决策变量。在这一表示之上，进一步建立客户端条件化效用、标准化训练成本、位置相关训练显存、计算时延和通信时延模型，使数据效用与系统资源可以在统一问题中进行权衡。

其次，研究将设计轻量的本地重要性画像。客户端加载当前全局适配器，从本地训练集中抽取一个小批次完成一次前向与反向传播，以梯度平方构造FIM近似分数，从而区分不同目标矩阵对当前客户端的潜在贡献。该过程不需要上传原始样本，画像结果还可在若干轮内缓存复用，以避免频繁评估带来的额外开销。

在获得矩阵效用后，研究将处理客户端内部组合空间。训练显存被分解为基础固定占用、随激活矩阵增加的参数与优化器状态，以及由最早激活位置决定的中间激活保留。客户端动态规划在显存和时延约束下保留各成本状态中效用较高的配置，并利用支配关系删除没有后续选择价值的候选，形成紧凑配置集合。

服务器接收各客户端的候选配置后，在共享轮次预算下联合决定客户端是否参与以及参与时采用哪一项配置。调度目标综合考虑画像效用和最慢客户端时延，使局部最优配置不会因争夺同一全局预算而破坏整体资源效率。完成调度后，客户端仅训练被分配的固定秩LoRA矩阵，并对返回参数执行联邦聚合。

实验部分将构建图像分类、自然语言理解和生成式摘要三类任务，设置IID、Dirichlet non-IID及长度偏斜数据划分，与固定配置、均匀矩阵、物理层分配、异构秩以及量化—秩联合分配方法比较。评价不仅考察最佳和最终任务指标，还考察累计预算—质量、时间—质量、配置可行性及关键模块消融，从多个角度验证方法的有效性。

课题拟重点解决四个相互关联的问题：如何用极少本地计算获得能够区分矩阵位置并反映non-IID客户端差异的效用信号；如何表达同数量但不同位置矩阵配置的训练显存差异；如何在不枚举指数规模子集的情况下保留优质候选；以及如何在共享预算、客户端个体约束和同步时延之间实现全局协调。"""

METHOD_ROUTE = """在系统建模阶段，以客户端、通信轮次和目标矩阵为基本索引，用二元向量表示每个客户端本轮激活的矩阵集合。配置效用由被激活矩阵的本地画像分数组成，标准化成本按固定秩rank-8矩阵等价数量计算。客户端受到训练显存、训练时延和激活矩阵数量约束，服务端受到轮次共享预算约束，优化目标是在提高聚合画像效用的同时抑制同步轮次时延。

本地矩阵画像由服务器下发的当前全局适配器触发。客户端在一个本地mini-batch上计算平均训练损失和梯度，对每个目标矩阵对应的LoRA参数求梯度平方和，得到FIM近似分数；随后在统一尺度上归一化为调度效用。画像按照固定间隔刷新，其余轮次复用缓存，使重要性信号能够随模型状态变化而更新，又不必在每轮重复完整计算。

训练显存被分解为基础固定占用、随激活矩阵数增长的参数与优化器状态，以及从最早激活矩阵到输出端需要保留的位置相关中间激活。空配置单独定义为固定占用，以避免最早激活位置未定义。这一模型能够区分“矩阵数相同但训练起点不同”的配置。

客户端按矩阵位置和成本建立动态规划状态，在每次转移时检查显存与时延可行性；对相同成本及最早位置状态只保留效用最大的配置，并删除被更低成本、更低时延且更高效用配置支配的项。该过程将客户端内部组合搜索与服务端选择解耦。

服务端对所有客户端菜单进行配置级协调。服务器枚举可行时延阈值，在预算内选择总效用最大的客户端—配置集合，再比较“效用减去时延惩罚”的目标，得到本轮调度结果。客户端只实例化并训练被选中的LoRA矩阵，服务端按参数键完成联邦平均。

实验采用CIFAR100与ViT-base、MNLI和QQP与DeBERTa-v3-large、QNLI和SST-2与RoBERTa-base、SAMSum与OPT-1.3B的组合。系统设置10个客户端以及分层训练显存、计算因子和带宽；分类任务报告Accuracy，摘要任务报告ROUGE-L，并与FedAvg-LoRA、Uniform、Fed-HeLLo、MS、LEGEND、HETLoRA-style和FAH-QLoRA-style进行比较。

可行性分析：项目已实现LoRA候选发现、本地FIM画像、客户端配置DP、服务端MCKP调度、联邦训练与评估管线；已有A800 GPU实验和完整结果台账。矩阵评分仅需一个本地批次，客户端画像和配置生成可并行执行；服务端只处理压缩后的菜单，避免全局穷举。数据集、模型和基线均有公开实现或论文依据，现有代码与实验基础能够支撑后续完善。"""

INNOVATION = """3. 创新之处

（1）提出矩阵级LoRA资源调度粒度。与按秩、量化精度或完整物理层分配不同，将具体目标矩阵身份作为独立决策变量，在相同训练容量下允许保留同层高效用投影并排除低效用投影。

（2）提出客户端条件化的轻量重要性画像。利用当前全局适配器和客户端一个本地小批次估计矩阵级FIM分数，不依赖服务端代理数据集，使调度信号能够反映non-IID客户端的局部数据差异。

（3）设计位置相关训练显存感知的两级调度。客户端DP显式考虑最早激活位置对中间激活保留的影响，生成满足显存和时延约束的紧凑菜单；服务端再以全局MCKP协调客户端与配置，在共享预算和同步时延下优化整体效用。

（4）建立跨视觉、语言理解和生成任务的统一验证框架，从峰值任务质量、累计预算和估计时间三个维度评价细粒度调度收益，并通过规范的实验台账保证结果可复核。"""

FOUNDATION = """1. 相关的研究工作积累

已完成FedMATRIX-LoRA原型系统：能够从不同Transformer主干动态发现LoRA候选矩阵；支持按客户端矩阵子集注入固定秩LoRA；实现单批次本地FIM画像、位置相关训练显存估计、客户端动态规划候选生成、服务端多选背包调度以及联邦聚合与逐轮评估。项目还实现FedAvg-LoRA、Uniform、Fed-HeLLo、MS、LEGEND、HETLoRA-style与FAH-QLoRA-style等对比机制。

已完成覆盖14个场景的主实验整理，包括CIFAR100三种划分、GLUE四项任务的IID与Dirichlet 1.0划分，以及SAMSum的IID、轻度和重度长度偏斜。现有结果覆盖图像分类、句子与句对理解、文本摘要三种类型，也包含从轻度到较强的数据异构设置。代表性结果中，本方法在CIFAR100 IID、MNLI IID和QNLI IID上的最佳Accuracy分别达到0.9000、0.8800和0.7960；更完整的运行标识、随机种子、配置文件和原始CSV由项目实验台账单独维护，不在开题报告正文中展开。

从目前的曲线和结果矩阵看，矩阵级选择不仅影响最终任务指标，也影响达到同一质量所需的累计训练容量和估计时间。这说明课题所研究的“容量放置位置”并非单纯实现细节，而是能够与客户端数据特征及资源约束共同作用的调度维度。现有实验还发现，图像、句子分类、句对推理和文本生成任务对矩阵位置的偏好并不完全相同，进一步支持采用客户端画像而不是固定全局排序。

已完成矩阵位置动机实验、矩阵粒度与物理层粒度对比、本地画像消融、深度校准消融、MCKP与贪心调度对比及延迟惩罚分析，并生成方法架构图、结果矩阵、预算曲线和时间曲线。上述工作为学位论文的系统实现、理论表述和实验章节提供了直接基础。

2. 已具备的研究条件，尚缺少的研究条件和拟解决的途径

已具备条件：PyTorch、Transformers、PEFT与联邦训练代码；CIFAR100、GLUE和SAMSum处理流程；A800 GPU与Slurm算力环境；模型缓存、实验清单、运行脚本、原始CSV、图表生成脚本和论文素材；Fed-HeLLo、LEGEND、HETLoRA与FAH-QLoRA等相关文献。

尚需完善：①选取代表性主场景补充多随机种子均值与方差，以增强统计稳健性；②用若干真实配置测量峰值训练显存和墙钟时间，校准解析资源模型；③进一步统一历史运行元数据并完成论文级复现实验入口；④在更大客户端规模下测试菜单规模与调度耗时。拟通过小规模针对性补跑、PyTorch显存插桩、统一manifest和自动化汇总脚本解决，不改变课题主线。

3. 本人已取得的主要成果

已形成可运行的矩阵级联邦LoRA调度代码、实验配置与结果管理流程；完成跨视觉、语言理解和文本生成任务的主实验及主要消融实验；完成英文论文初稿、中文说明材料、方法图与实验图表。论文投稿、软件著作权或其他正式成果如尚未获得，将在最终提交时按实际情况填写，不作预先认定。"""

EXPECTED = """预期成果

1. 完成一套面向异构客户端的矩阵级LoRA联邦微调方法与可复现实验系统。
2. 形成训练显存、计算/通信时延与轮次预算约束下的数学模型、客户端候选配置算法和服务端联合调度算法。
3. 完成覆盖视觉分类、语言理解和文本摘要任务的实验评估与消融分析，形成可追溯数据表和图表。
4. 完成研究生学位论文；在条件允许的情况下形成高水平学术论文投稿和开源代码材料。
"""


REVIEW_ABSTRACT = (
    "基础模型在边缘设备和跨机构场景中的落地，使隐私保护与资源受限条件下的协同微调成为重要问题。联邦学习能够避免集中原始数据，参数高效微调则通过仅训练少量参数降低计算与通信开销，两者结合形成联邦参数高效微调。本文围绕异构联邦LoRA的资源分配研究进行综述。首先梳理联邦学习、Transformer、LoRA及量化微调的发展；随后从联邦PEFT与聚合、异构秩和量化、重要性引导的层分配、资源感知客户端调度四个方向分析代表性工作；进一步比较秩、层、矩阵和量化等适配维度，以及服务端统一评分与客户端本地评分的差异。分析表明，现有研究已能降低可训练参数并适应部分设备异构，但对同一物理层内不同目标矩阵的效用差异、非独立同分布数据下的客户端特定重要性，以及矩阵位置引起的训练显存差异关注不足。未来值得研究客户端本地轻量画像、位置相关资源建模、候选配置压缩与全局联合调度，并加强真实设备测量、统计稳健性和大规模扩展验证。"
)

REVIEW_SECTIONS = [
    ("一、前言", [
        "大规模预训练模型依靠通用语料和图像学习可迁移表示，已经成为自然语言处理、计算机视觉和生成式人工智能的重要基础。ViT、RoBERTa、DeBERTa和OPT等模型分别展示了Transformer在视觉识别、语言理解和文本生成中的迁移能力[34-37]。然而，通用预训练并不能直接覆盖医疗、金融、通信、政务和个人终端中的具体业务需求，模型仍需利用领域数据进行适配。这些数据往往分散在不同组织或设备中，并受到隐私、所有权和合规要求限制，难以集中到同一训练平台。",
        "联邦学习为这一矛盾提供了基本解决思路。FedAvg允许客户端在本地执行若干步优化，只向服务器上传模型更新并完成全局聚合[1]。随后，大量研究从统计异构、设备异构、隐私保护和通信效率等方面扩展这一范式[2-4]。联邦学习由此能够在不直接汇集原始数据的前提下利用多方知识，但当参与训练的对象从中小型网络转向基础模型时，完整模型的训练显存、计算量与通信量迅速成为新的瓶颈。",
        "参数高效微调正是在这一背景下进入联邦场景。Adapter通过插入小型瓶颈模块减少可训练参数[6]，Prefix-Tuning、P-Tuning和BitFit分别从连续提示或偏置参数入手完成任务适配[7-9]。LoRA则冻结原始权重，用两个低秩矩阵表示线性层的增量，在保持较强适配能力的同时显著降低可训练参数和通信负载[5]；QLoRA进一步量化冻结主干，使更大的模型能够在有限显存中完成微调[10]。联邦学习解决数据不能集中，PEFT解决完整模型难以训练，两条技术路线的结合自然形成联邦参数高效微调。",
        "不过，参数数量减少并不意味着所有客户端都应使用相同配置。客户端的本地数据可能呈现明显的non-IID分布，训练显存、算力和带宽也存在差异。统一LoRA配置可能使低资源客户端无法参与，也可能让高资源客户端的能力得不到利用；统一的重要性排序还可能忽略不同客户端对矩阵位置的不同需求。因此，联邦LoRA研究逐渐从“如何把LoRA放入联邦训练”转向“如何根据数据和资源为客户端分配LoRA容量”。本文沿着这一演进脉络，综述统一联邦PEFT、异构秩与量化、重要性引导的层分配以及资源感知调度，并进一步讨论矩阵级配置所提供的新研究空间。",
    ]),
    ("二、从联邦学习到异构联邦LoRA的技术演进", [
        "联邦学习早期工作的核心是减少分布式优化中的通信轮次。FedAvg通过增加本地训练步数，用较少的全局同步换取较多的客户端计算[1]。随着研究走向真实边缘环境，问题的重点逐渐从通信次数扩展到“哪些客户端适合参加本轮训练”。Oort综合客户端数据效用和运行速度选择参与者[28]，FedBalancer协调样本难度与训练节奏[29]，PyramidFL进一步进行细粒度客户端选择[30]，AutoFL和Cost-Effective FL则把能耗、计算与无线资源纳入决策[31-32]。这些研究共同说明：在异构系统中，仅规定统一训练流程是不够的，服务器需要根据客户端状态分配有限资源。",
        "与此同时，模型异构研究开始改变“所有客户端必须训练同一个完整模型”的假设。HeteroFL让不同能力的设备训练不同宽度的子网络[26]，FedRolex通过滚动抽取子模型改善参数覆盖[27]。它们虽然不是专门为LoRA设计，却建立了一个重要认识：模型内部的参数覆盖也可以成为联邦资源分配变量。客户端选择决定由谁训练，模型异构进一步决定每个客户端训练模型的哪一部分。",
        "PEFT的发展为这种结构化分配提供了更轻量的载体。与Adapter和提示调优相比，LoRA直接作用于Transformer线性变换，通常只需保存和传输少量低秩因子[5-9]。FedPETuning系统研究了多类PEFT方法在联邦环境中的表现[11]，FedTune和自适应提示优化工作进一步讨论了训练效率与优化策略[12-13]。这些研究首先证明了联邦PEFT的可行性，但大多采用统一的目标模块和适配容量，客户端之间的差异主要体现为本地数据而非LoRA结构。",
        "当统一配置的限制逐渐显现后，研究开始沿不同维度构造异构LoRA。HetLoRA允许客户端采用不同秩并设计相应聚合机制[16]；FlexLoRA利用服务器端合成与奇异值分解处理客户端任务和资源异构[17]；LEGEND依据设备能力和参数重要性分配训练容量[19]；FAH-QLoRA把冻结主干的量化精度与LoRA秩联合起来[18]。另一条路线是改变适配位置，Fed-HeLLo根据层重要性和设备显存分配不同LoRA物理层[15]。至此，联邦LoRA的研究对象从单一聚合流程扩展为秩、精度、层覆盖和客户端参与的联合配置。",
        "这条演进路线还带来一个新的系统问题。秩、量化和层覆盖不再只是训练开始前固定的超参数，而可能随客户端资源和模型状态发生变化。服务器不仅需要合并训练后的更新，还要在训练前确定参与者和配置；客户端也需要判断候选配置是否满足本地约束。联邦LoRA因此逐渐具备调度问题的特征，其研究重点由单纯减少参数转向在任务效用、训练显存、时延和通信之间寻找更合适的资源—质量权衡。",
    ]),
    ("三、异构联邦LoRA的主要研究路线", [
        "现有研究首先关注LoRA在联邦聚合中的一致性。低秩更新由两个因子矩阵相乘得到，分别平均因子并不总是等价于平均完整权重增量。FedEx-LoRA围绕这一问题构造精确聚合方法[14]，LoRA-FAIR从初始化和聚合偏差出发提高训练稳定性[21]，Tensor-aggregated LoRA尝试在张量聚合前对齐异构适配器[22]，后续工作还结合流水线、误差抑制和矩阵冻结改善执行效率[23]。这一方向解决的是客户端完成本地训练后如何正确合并更新，它构成联邦LoRA的重要基础，但没有直接回答训练开始前应为客户端分配哪些适配参数。",
        "围绕训练前配置，异构秩是最早形成系统研究的一条路线。LoRA秩同时影响表达能力、参数数量、优化器状态和通信量，因此可以作为设备能力与训练容量之间的调节旋钮。HetLoRA、FlexLoRA、LEGEND和LoRA-A2分别采用异构秩、SVD重构、重要性分配或自适应冻结等机制[16-17,19-20]。它们使低资源客户端使用较小适配器，高资源客户端承担更多训练容量，比统一秩更符合设备异构环境。不过，秩主要回答每个更新具有多大容量；如果所有客户端仍训练固定位置，同样的参数预算可能被放在任务收益不同的矩阵上。",
        "量化研究从基础模型常驻存储的角度补充了秩分配。QLoRA表明，冻结主干可以低比特存储，而LoRA参数保持较高精度完成训练[10]。FAH-QLoRA将这一思想带入联邦环境，根据客户端能力联合确定量化精度和LoRA秩[18]。这种设计能够显著扩大可参与设备范围，并在主干存储与适配容量之间进行权衡。与此同时，量化精度与秩仍主要控制“使用多少资源”，尚未直接利用Q、K、V或前馈投影之间的功能差异。",
        "Fed-HeLLo把研究进一步推进到适配位置。该方法先在服务器代理数据上评估Transformer层重要性，再结合客户端显存容量决定可训练的LoRA层[15]。这一设计的重要意义在于承认不同位置的训练价值并不相同，资源有限时应优先选择更有贡献的部分。不过，物理层通常同时包含查询、键、值、输出投影和前馈网络等多个线性矩阵，整层分配会把这些矩阵绑定在一起。即使某个投影具有较高效用，客户端也可能必须同时承担同层其他矩阵的成本。",
        "重要性信号由谁计算同样影响方法适用范围。服务端统一评分实现简单，也便于向不同容量客户端映射同一排序，但它要求服务器拥有可用于评估的公共或代理数据。联邦客户端的数据往往具有non-IID特征[2,33]，一个统一代理集难以同时代表所有客户端。若把评分移到客户端本地，重要性可以同时依赖局部数据和当前全局适配器，更符合数据不上传的联邦设定；相应地，系统需要控制本地画像的计算频率、分数噪声以及额外的信息交换。",
        "从系统视角看，上述方法与资源感知联邦调度正在逐步汇合。传统调度主要决定客户端是否参与，并把本地训练配置视为已知[28-32]；异构LoRA则使客户端拥有多个精度、秩或位置配置。于是服务器面对的不再是单纯的客户端选择，而是客户端—配置联合选择。聚合发生在训练之后，调度发生在训练之前，两者可以使用同一联邦框架，但应当在问题定义和实验中清楚区分。",
    ]),
    ("四、从现有方法到矩阵级调度的研究空间", [
        "将现有方法放在统一视角下，可以发现秩、量化、物理层和矩阵分别描述了不同的适配维度。秩决定一个低秩更新具有多少表达容量，量化决定主干或通信表示使用多少位宽，物理层决定哪些Transformer块参与训练，而矩阵粒度进一步区分同一层内的Q、K、V、输出和前馈投影。前两类方法更接近“分配多少容量”，后两类方法则进一步回答“把容量放在哪里”。矩阵级决策能够表示任意物理层组合，同时允许在同一层内保留高效用投影并排除低效用投影，因此在预算紧张时提供了更细的选择空间。",
        "更细粒度并不自动带来更好结果，关键在于效用估计是否与客户端真正相关。服务器代理数据得到的全局排序适合表达总体趋势，却可能掩盖non-IID客户端之间的偏好差异。客户端本地评分能够直接利用私有数据，但若每轮遍历完整数据集，其开销可能抵消调度收益。一个自然折中是利用当前全局适配器，在少量本地样本上计算梯度平方或FIM近似，并以较低频率刷新。这样得到的信号不是训练后的准确率，而是配置搜索所需的轻量效用代理。",
        "矩阵位置还会改变资源成本的结构。LoRA参数本身相对较少，但客户端训练时仍需保存冻结主干、任务头、梯度、优化器状态和中间激活。若最早被训练的位置靠近模型输入端，反向传播需要经过更长路径并保留更多后续激活；若只训练靠近输出端的矩阵，激活路径通常更短。因此，相同矩阵数或相同秩不一定对应相同训练显存。仅以参数量或层数估算资源，可能把在真实训练中不可行的配置交给低显存客户端。",
        "从矩阵级效用与位置相关成本出发，配置搜索会迅速形成组合爆炸。具有L个候选矩阵时，潜在子集数量随L指数增长；多个客户端联合后更无法由服务器直接枚举。适合的解决方式是分层处理：每个客户端先根据本地效用、显存和时延生成少量非支配配置，服务器再把这些配置作为菜单进行全局选择。客户端阶段保留个体可行性，服务器阶段协调共享预算与同步时延，从而在不穷举全部组合的情况下完成联合调度。",
        "这种研究路线也要求更完整的评价方法。最终准确率或ROUGE只能说明训练结果，不能回答达到结果付出了多少资源。异构联邦PEFT应同时观察最佳和最终任务指标、达到目标质量所需的累计训练容量、估计或实测时间、显存可行性与通信量。时间横轴若来自解析模型，就应与真实墙钟测量区分；预算若采用标准化矩阵成本，也需要说明它与参数量、FLOPs和实际训练时间的关系。",
        "公平比较同样是该方向的难点。不同方法可能同时改变秩、目标层、客户端数量、训练步数和量化精度，只固定最终参数量并不能保证反向传播路径与优化难度一致。较可靠的实验应统一模型、数据划分、局部优化器、轮数和评价流程，再让各方法在相同资源口径下作出配置决策。对于依据论文核心思想实现的机制级基线，还应明确保留了哪些决策，避免将简化版本描述为原系统的完整复现。",
        "最后，本地画像虽然不需要上传原始样本，但画像分数仍是由本地数据导出的统计量。当前联邦学习通常允许交换模型更新和训练统计，未来仍可进一步结合安全聚合、差分隐私、候选菜单上传或排序压缩，减少完整分数暴露。矩阵级调度由此不仅是优化问题，也与隐私、系统实现和可解释评价紧密相关。",
    ]),
    ("五、研究趋势与未来方向", [
        "未来的联邦LoRA配置很可能由单一维度走向多维联合决策。矩阵位置决定容量放在哪里，秩决定每个位置分配多少容量，量化决定主干存储精度，客户端参与和本地步数又决定本轮训练范围。把所有变量一次性交给统一优化器会导致搜索空间急剧膨胀，因此更现实的方向是建立分层决策：先由设备约束排除不可行区域，再由本地画像生成候选菜单，最后由服务器协调全局预算。Pareto剪枝、近似动态规划和分布式候选生成将在其中发挥重要作用。",
        "随着全局适配器持续更新，矩阵重要性不会始终保持不变；客户端的带宽、可用显存和计算负载也可能随时间波动。静态配置因而会逐步让位于动态但低开销的重配置机制。低频画像、缓存复用、指数滑动平均和事件触发更新能够在信号新鲜度与额外开销之间取得平衡。进一步结合在线资源预测后，调度器可以从响应当前状态发展为提前规避慢客户端和资源峰值。",
        "评价平台也将由单纯解析仿真走向真实设备与仿真结合。解析模型适合大规模客户端实验，但训练显存峰值、GPU并行行为、能耗和网络抖动最终仍需通过设备测量校准。较理想的研究范式是先在真实GPU上测量代表性配置，再用拟合后的资源模型扩展到更大的客户端群体，并报告估计误差。这样既保留实验规模，又避免把模型时间误解为真实墙钟时间。",
        "跨任务迁移将成为检验算法通用性的另一条主线。视觉Transformer、编码器模型和生成式解码器在目标矩阵功能、序列长度与激活结构上存在差异，某一任务上有效的固定深度或投影偏好未必能直接迁移。相比手工规定统一位置，客户端条件化效用与显式资源约束提供了更通用的适配方式。未来还可以研究历史任务形成的先验如何与当前本地画像结合，在减少探测开销的同时保持任务自适应能力。",
        "理论分析将更多围绕系统行为的可解释性展开。候选菜单是否保留各成本下的优质配置、调度结果是否满足硬约束、算法复杂度如何随客户端和预算增长、资源估计误差是否会改变配置可行性，这些问题直接决定方法能否部署和扩展。与实现路径一致的性质证明和复杂度分析，比脱离系统假设的抽象结论更能支撑联邦微调调度研究。",
        "隐私与鲁棒性也会伴随细粒度画像进入研究重点。局部梯度统计可能受到恶意客户端操纵，画像上传也可能泄露局部数据特征。安全聚合、鲁棒归一化、异常分数检测和隐私预算控制可以与资源调度结合，使系统在追求效率的同时保持联邦学习的信任边界。",
    ]),
    ("六、总结与展望", [
        "回顾相关研究可以看到，联邦学习、参数高效微调和异构资源调度并不是彼此孤立的方向。联邦学习首先解决数据不可集中问题，PEFT使基础模型能够以较少可训练参数参与协同训练，异构LoRA又进一步让不同能力的客户端使用不同适配容量。研究由统一Adapter或LoRA逐步发展到异构秩、联合量化和重要性层分配，反映出一个持续变化的核心问题：如何把有限训练资源放到真正能够产生任务收益的位置。",
        "现有方法已经证明结构异构能够提升设备包容性和资源效率，但客户端局部数据、具体目标矩阵身份和位置相关训练显存仍未被充分统一。矩阵级调度为此提供了一条自然延伸路线。它不必改变LoRA的基本低秩结构，而是把目标矩阵身份视为可调度资源；本地轻量画像负责表达non-IID客户端的效用差异，位置相关资源模型负责判断配置是否可训练，客户端候选压缩和服务端联合选择则负责控制组合复杂度并协调共享预算。",
        "进一步研究应在算法收益和证据质量两方面继续完善。一方面，可以探索矩阵、秩与量化的多维配置，以及更稳健、更隐私友好的动态画像；另一方面，应加强多随机种子实验、真实设备显存与墙钟时间测量、大规模客户端扩展和统一预算口径。只有让问题建模、代码实现、资源度量和实验结论保持一致，细粒度联邦LoRA调度才能从有效的实验方案发展为可解释、可复现并具有部署价值的系统方法。",
    ]),
]


REFERENCES = [
    "[1] McMahan B, Moore E, Ramage D, et al. Communication-efficient learning of deep networks from decentralized data[C]//AISTATS. 2017: 1273-1282.",
    "[2] Kairouz P, McMahan H B, Avent B, et al. Advances and open problems in federated learning[J]. Foundations and Trends in Machine Learning, 2021, 14(1-2): 1-210.",
    "[3] Li T, Sahu A K, Talwalkar A, et al. Federated learning: Challenges, methods, and future directions[J]. IEEE Signal Processing Magazine, 2020, 37(3): 50-60.",
    "[4] Lim W Y B, Luong N C, Hoang D T, et al. Federated learning in mobile edge networks: A comprehensive survey[J]. IEEE Communications Surveys & Tutorials, 2020, 22(3): 2031-2063.",
    "[5] Hu E J, Shen Y, Wallis P, et al. LoRA: Low-rank adaptation of large language models[C]//ICLR. 2022.",
    "[6] Houlsby N, Giurgiu A, Jastrzebski S, et al. Parameter-efficient transfer learning for NLP[C]//ICML. 2019: 2790-2799.",
    "[7] Li X L, Liang P. Prefix-tuning: Optimizing continuous prompts for generation[C]//ACL-IJCNLP. 2021: 4582-4597.",
    "[8] Liu X, Ji K, Fu Y, et al. P-Tuning: Prompt tuning can be comparable to fine-tuning across scales[C]//ACL. 2022: 61-68.",
    "[9] Ben Zaken E, Goldberg Y, Ravfogel S. BitFit: Simple parameter-efficient fine-tuning for transformer-based masked language models[C]//ACL. 2022: 1-9.",
    "[10] Dettmers T, Pagnoni A, Holtzman A, et al. QLoRA: Efficient finetuning of quantized LLMs[C]//NeurIPS. 2023, 36: 10088-10115.",
    "[11] Zhang Z, Yang Y, Dai Y, et al. FedPETuning: When federated learning meets the parameter-efficient tuning methods of pre-trained language models[C]//Findings of ACL. 2023: 9963-9977.",
    "[12] Chen J, Xu W, Guo S, et al. FedTune: A deep dive into efficient federated fine-tuning with pre-trained transformers[EB/OL]. arXiv:2211.08025, 2022.",
    "[13] Che T, Liu J, Zhou Y, et al. Federated learning of large language models with parameter-efficient prompt tuning and adaptive optimization[C]//EMNLP. 2023: 7871-7888.",
    "[14] Singhal R, Ponkshe K, Vepakomma P. FedEx-LoRA: Exact aggregation for federated and efficient fine-tuning of large language models[C]//ACL. 2025: 1316-1336.",
    "[15] Zhang Z, Liu P, Xu J, et al. Fed-HeLLo: Efficient federated foundation model fine-tuning with heterogeneous LoRA allocation[J]. IEEE Transactions on Neural Networks and Learning Systems, 2025, 36(10): 17556-17569.",
    "[16] Cho Y J, Liu L, Xu Z, et al. Heterogeneous LoRA for federated fine-tuning of on-device foundation models[C]//EMNLP. 2024: 12903-12913.",
    "[17] Bai J, Chen D, Qian B, et al. Federated fine-tuning of large language models under heterogeneous tasks and client resources[C]//NeurIPS. 2024, 37: 14457-14483.",
    "[18] Gao Z, Zhang Z, Guo Y, et al. Federated adaptive fine-tuning of large language models with heterogeneous quantization and LoRA[C]//IEEE INFOCOM. 2025: 1-10.",
    "[19] Liu J, Liao Y, Xu H, et al. Adaptive parameter-efficient federated fine-tuning on heterogeneous devices[J]. IEEE Transactions on Mobile Computing, 2025, 24(11): 12533-12549.",
    "[20] Koo J, Jang M, Ok J. Towards robust and efficient federated low-rank adaptation with heterogeneous clients[C]//ACL. 2025: 416-429.",
    "[21] Bian J, Wang L, Zhang L, et al. LoRA-FAIR: Federated LoRA fine-tuning with aggregation and initialization refinement[C]//ICCV. 2025: 3737-3746.",
    "[22] Li Z, Xu B, Shu X, et al. Tensor-aggregated LoRA in federated fine-tuning[C]//ICCV. 2025: 1058-1067.",
    "[23] Wang H, Wang X, Li Y, et al. Federated LoRA fine-tuning with pipelined error-mitigated aggregation and matrix-wise freezing[C]//Findings of ACL. 2026: 5749-5762.",
    "[24] Huang X, Hu Y, Gong Y J, et al. GMFL: Efficient global masking for federated LLM fine-tuning[C]//ACL. 2026: 25293-25312.",
    "[25] Ban H, Ji K. Rethinking parameter sharing for LLM fine-tuning with multiple LoRAs[C]//Findings of ACL. 2026: 12843-12863.",
    "[26] Diao E, Ding J, Tarokh V. HeteroFL: Computation and communication efficient federated learning for heterogeneous clients[C]//ICLR. 2021.",
    "[27] Alam S, Liu L, Yan M, et al. FedRolex: Model-heterogeneous federated learning with rolling sub-model extraction[C]//NeurIPS. 2022, 35: 29677-29690.",
    "[28] Lai F, Zhu X, Madhyastha H V, et al. Oort: Efficient federated learning via guided participant selection[C]//OSDI. 2021: 19-35.",
    "[29] Shin J, Li Y, Liu Y, et al. FedBalancer: Data and pace control for efficient federated learning on heterogeneous clients[C]//MobiSys. 2022: 436-449.",
    "[30] Li C, Zeng X, Zhang M, et al. PyramidFL: A fine-grained client selection framework for efficient federated learning[C]//MobiCom. 2022: 158-171.",
    "[31] Kim Y G, Wu C J. AutoFL: Enabling heterogeneity-aware energy efficient federated learning[C]//MICRO. 2021: 183-198.",
    "[32] Luo B, Li X, Wang S, et al. Cost-effective federated learning in mobile edge networks[J]. IEEE Journal on Selected Areas in Communications, 2021, 39(12): 3606-3621.",
    "[33] Li Q, He B, Song D. Model-contrastive federated learning[C]//CVPR. 2021: 10713-10722.",
    "[34] Dosovitskiy A, Beyer L, Kolesnikov A, et al. An image is worth 16x16 words: Transformers for image recognition at scale[C]//ICLR. 2021.",
    "[35] Liu Y, Ott M, Goyal N, et al. RoBERTa: A robustly optimized BERT pretraining approach[EB/OL]. arXiv:1907.11692, 2019.",
    "[36] He P, Liu X, Gao J, et al. DeBERTa: Decoding-enhanced BERT with disentangled attention[C]//ICLR. 2021.",
    "[37] Zhang S, Roller S, Goyal N, et al. OPT: Open pre-trained transformer language models[EB/OL]. arXiv:2205.01068, 2022.",
    "[38] Wang A, Singh A, Michael J, et al. GLUE: A multi-task benchmark and analysis platform for natural language understanding[C]//ICLR. 2019.",
    "[39] Gliwa B, Mochol I, Biesek M, et al. SAMSum corpus: A human-annotated dialogue dataset for abstractive summarization[C]//New Frontiers in Summarization. 2019: 70-79.",
]


def set_east_asia_font(run, name="宋体", size=12, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def format_paragraph(p, size=12, bold=False, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    for run in p.runs:
        set_east_asia_font(run, size=size, bold=bold)


def add_text_block(container, text):
    container.text = ""
    append_text_block(container, text, reuse_first=True)


def append_text_block(container, text, reuse_first=False):
    for idx, block in enumerate(text.split("\n\n")):
        p = container.paragraphs[0] if reuse_first and idx == 0 else container.add_paragraph()
        p.add_run(block)
        is_heading = len(block) < 55 and (
            block[:2].rstrip("、.") in {"1", "2", "3"} or block.startswith("（")
        )
        format_paragraph(p, size=12, bold=is_heading, indent=not is_heading)


def add_reference_list(container, refs, size=9.5):
    for ref in refs:
        p = container.add_paragraph(ref)
        format_paragraph(p, size=size, indent=False)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.space_after = Pt(1)


def add_omml_equation(container, text):
    """Append a centered Office Math object using readable Unicode math."""
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    math_pr = OxmlElement("m:oMathParaPr")
    jc = OxmlElement("m:jc")
    jc.set(qn("m:val"), "center")
    math_pr.append(jc)
    math_para.append(math_pr)
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    run_pr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    run_pr.append(sty)
    run.append(run_pr)
    node = OxmlElement("m:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    math.append(run)
    math_para.append(math)
    p._p.append(math_para)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)


def delete_table_row(table, row_index=0):
    table._tbl.remove(table.rows[row_index]._tr)


def replace_paragraph_text(p, text, size=14, bold=False, center=False):
    p.clear()
    r = p.add_run(text)
    set_east_asia_font(r, size=size, bold=bold)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_cell_picture(cell, path, width_cm, caption):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = cell.add_paragraph(caption)
    format_paragraph(cap, size=10.5, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)


def remove_element(el):
    el.getparent().remove(el)


def generate_proposal():
    dst = OUT / "研究生学位论文开题报告_FedMATRIX-LoRA.docx"
    shutil.copy2(PROPOSAL_TEMPLATE, dst)
    doc = Document(dst)

    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("学位（毕业）论文题目"):
            replace_paragraph_text(p, f"学位（毕业）论文题目：{ZH_TITLE}", 18, True, True)
        elif t == "二〇二六年四月":
            replace_paragraph_text(p, "二〇二六年十月", 16, False, True)
        elif t == "博士研究生               硕士研究生":
            replace_paragraph_text(p, "博士研究生              ☑ 硕士研究生", 14, False, True)

    t0 = doc.tables[0]
    t0.cell(0, 2).text = ZH_TITLE
    t0.cell(1, 2).text = EN_TITLE
    t0.cell(2, 2).text = "□ 是\n☑ 否，并承诺学位论文将不涉及国家秘密和其他不宜公开的内容。"
    add_text_block(t0.cell(3, 2), PROPOSAL_ABSTRACT)
    t0.cell(4, 2).text = "联邦学习  参数高效微调  低秩适配  矩阵级调度  资源异构"
    t0.cell(5, 2).text = "Federated Learning / Parameter-Efficient Fine-Tuning / Low-Rank Adaptation / Matrix-Level Scheduling / Resource Heterogeneity"
    for row in t0.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                format_paragraph(p, size=11, indent=False)
    delete_table_row(t0, 6)

    basis_cell = doc.tables[1].cell(1, 0)
    add_text_block(basis_cell, SIGNIFICANCE + "\n\n" + STATE_OF_ART + "\n\n3、主要参考文献及出处")
    add_reference_list(basis_cell, REFERENCES[:25], size=9.5)
    delete_table_row(doc.tables[1], 0)

    plan_cell = doc.tables[2].cell(1, 0)
    add_text_block(plan_cell, RESEARCH_PLAN)
    add_cell_picture(
        plan_cell,
        ROOT / "paper_assets/05_figures/fig_structure/preview/fig2_text_aligned_final.png",
        15.0,
        "图1  FedMATRIX-LoRA总体技术路线",
    )
    method_parts = METHOD_ROUTE.split("\n\n")
    append_text_block(plan_cell, method_parts[0])
    p = plan_cell.add_paragraph(
        "为明确矩阵级决策，设 i 表示客户端，k 表示通信轮次，j 表示候选矩阵位置，L 表示模型中可注入LoRA的候选矩阵总数。客户端 i 在第 k 轮的激活向量定义为："
    )
    format_paragraph(p, size=12, indent=True)
    add_omml_equation(plan_cell, "𝐦ᵢᵏ = [mᵢ,₁ᵏ, mᵢ,₂ᵏ, …, mᵢ,ᴸᵏ],   mᵢ,ⱼᵏ ∈ {0, 1}")
    p = plan_cell.add_paragraph(
        "其中，m 的取值为1表示对应目标矩阵在本轮被激活并参与本地训练，取值为0表示不训练。该向量同时决定配置效用、矩阵成本和训练资源消耗。"
    )
    format_paragraph(p, size=12, indent=True)

    append_text_block(plan_cell, method_parts[1])
    p = plan_cell.add_paragraph(
        "设 θⱼ 表示候选矩阵 j 对应的LoRA参数集合，ℒᵢ表示客户端 i 在画像小批次上的平均训练损失。矩阵的原始FIM近似分数及其归一化效用分别写为："
    )
    format_paragraph(p, size=12, indent=True)
    add_omml_equation(plan_cell, "gᵢ,ⱼᵏ = ∑ₚ∈θⱼ (∂ℒᵢ / ∂p)²")
    add_omml_equation(plan_cell, "vᵢ,ⱼᵏ = gᵢ,ⱼᵏ / maxₕ,ₗ gₕ,ₗᵏ")
    p = plan_cell.add_paragraph(
        "这里，g 是梯度平方和得到的原始分数；v 是映射到统一尺度后的矩阵效用；h 和 l 分别遍历参与画像的客户端和候选矩阵。一个配置的画像效用采用其所有激活矩阵效用之和："
    )
    format_paragraph(p, size=12, indent=True)
    add_omml_equation(plan_cell, "Vᵢᵏ(𝐦ᵢᵏ) = ∑ⱼ₌₁ᴸ mᵢ,ⱼᵏ vᵢ,ⱼᵏ")

    append_text_block(plan_cell, method_parts[2])
    p = plan_cell.add_paragraph(
        "对任一非空配置，训练显存估计采用下式。它不只统计LoRA参数，而是同时考虑冻结模型等固定占用、每个激活矩阵增加的训练状态，以及由反向传播起点决定的中间激活："
    )
    format_paragraph(p, size=12, indent=True)
    add_omml_equation(plan_cell, "Mᵢ(𝐦) = Mᵢᶠⁱˣᵉᵈ + qᵢ(r)c(𝐦) + ∑ⱼ₌ⱼₘᵢₙᴸ μᵢ,ⱼ")
    p = plan_cell.add_paragraph(
        "其中，Mᶠⁱˣᵉᵈ表示客户端基础训练占用，qᵢ(r)表示每激活一个秩为 r 的LoRA矩阵所需的参数、梯度和优化器状态，c(𝐦)表示激活矩阵数，jₘᵢₙ表示最早被激活的位置，μᵢ,ⱼ表示位置 j 对应的中间激活占用。矩阵位置越靠前，需要保留的后续激活通常越多。"
    )
    format_paragraph(p, size=12, indent=True)

    append_text_block(plan_cell, method_parts[3])
    append_text_block(plan_cell, method_parts[4])
    p = plan_cell.add_paragraph(
        "服务器的目标是在本轮预算和各客户端可行性约束下，提高所有被选配置的总画像效用，同时避免最慢客户端造成过大的同步等待："
    )
    format_paragraph(p, size=12, indent=True)
    add_omml_equation(plan_cell, "max  ∑ᵢ Vᵢᵏ(𝐦ᵢᵏ) − λτᵏ")
    p = plan_cell.add_paragraph(
        "其中，τᵏ表示第 k 轮所有参与客户端中的最大训练时延，λ是效用与时延之间的权衡系数。服务器并不直接枚举所有矩阵子集，而是在客户端生成的紧凑配置菜单上完成联合选择。"
    )
    format_paragraph(p, size=12, indent=True)

    append_text_block(plan_cell, "\n\n".join(method_parts[5:]))
    append_text_block(plan_cell, INNOVATION)
    delete_table_row(doc.tables[2], 0)

    foundation_cell = doc.tables[3].cell(1, 0)
    add_text_block(foundation_cell, FOUNDATION)
    add_cell_picture(
        foundation_cell,
        ROOT / "paper_assets/05_figures/final/fig1_main_result_matrix.png",
        15.0,
        "图2  已完成实验的多任务最佳指标矩阵",
    )
    add_cell_picture(
        foundation_cell,
        ROOT / "paper_assets/05_figures/final/app_full_budget_curves.png",
        15.0,
        "图3  分类任务中的累计预算—性能曲线",
    )
    add_cell_picture(
        foundation_cell,
        ROOT / "paper_assets/05_figures/final/app_full_time_curves.png",
        15.0,
        "图4  分类任务中的累计估计时间—性能曲线",
    )
    delete_table_row(doc.tables[3], 0)

    schedule = [
        ("2026.10—2026.11", "完善文献调研、问题建模与开题论证", "完成开题"),
        ("2026.11—2027.01", "统一代码、配置清单与运行元数据，补充关键多种子实验", "阶段实验"),
        ("2027.01—2027.02", "开展真实显存与墙钟时间插桩，校准资源模型", "模型校准"),
        ("2027.02—2027.03", "完成大规模与关键消融实验，整理统计结果", "实验定稿"),
        ("2027.03—2027.04", "撰写和修改学位论文方法、系统与实验章节", "论文初稿"),
        ("2027.04—2027.05", "根据导师意见修改论文，完成图表、参考文献与查重", "论文定稿"),
        ("2027.05—2027.06", "准备送审、答辩材料和代码归档", "答辩准备"),
    ]
    table = doc.tables[4]
    for idx, (period, work, note) in enumerate(schedule, 2):
        table.cell(idx, 0).text = period
        table.cell(idx, 1).text = work
        table.cell(idx, 2).text = note
    add_text_block(table.cell(11, 0), EXPECTED)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                format_paragraph(p, size=10.5, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER if len(p.text) < 30 else WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.save(dst)
    return dst


def generate_review():
    dst = OUT / "研究生学位论文文献综述_异构联邦LoRA.docx"
    shutil.copy2(REVIEW_TEMPLATE, dst)
    doc = Document(dst)
    body = doc._element.body
    # Keep the official cover and signature line, remove template instructions.
    for p in list(doc.paragraphs):
        if p.text.strip() == "说  明":
            remove = False
            for el in list(body):
                if el is p._element:
                    remove = True
                if remove and el.tag != qn("w:sectPr"):
                    remove_element(el)
            break

    for p in doc.paragraphs:
        if p.text.strip() == "二〇二六年四月":
            replace_paragraph_text(p, "二〇二六年十月", 16, False, True)
        elif p.text.strip() == "博士研究生               硕士研究生":
            replace_paragraph_text(p, "博士研究生              ☑ 硕士研究生", 14, False, True)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    p.add_run("异构联邦微调中的LoRA资源分配与矩阵级调度研究综述")
    format_paragraph(p, size=18, bold=True, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph()
    p.add_run("研究生姓名：________________")
    format_paragraph(p, size=14, bold=True, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)

    p = doc.add_paragraph()
    r = p.add_run("摘要：")
    set_east_asia_font(r, "黑体", 12, True)
    r = p.add_run(REVIEW_ABSTRACT)
    set_east_asia_font(r, "宋体", 12, False)
    format_paragraph(p, size=12, indent=False)
    p = doc.add_paragraph()
    r = p.add_run("关键词：")
    set_east_asia_font(r, "黑体", 12, True)
    r = p.add_run("联邦学习；参数高效微调；LoRA；异构资源；矩阵级调度")
    set_east_asia_font(r, "宋体", 12, False)
    format_paragraph(p, size=12, indent=False)

    for heading, paragraphs in REVIEW_SECTIONS:
        hp = doc.add_paragraph()
        hp.add_run(heading)
        format_paragraph(hp, size=12, bold=True, indent=False)
        for text in paragraphs:
            p = doc.add_paragraph(text)
            format_paragraph(p, size=12, indent=True)

    hp = doc.add_paragraph("参考文献")
    format_paragraph(hp, size=12, bold=True, indent=False)
    for ref in REFERENCES:
        p = doc.add_paragraph(ref)
        format_paragraph(p, size=10.5, indent=False)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)

    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.5)
    doc.save(dst)
    return dst


def latex_escape(s):
    repl = {
        "&": r"\&", "%": r"\%", "#": r"\#", "_": r"\_\allowbreak{}", "$": r"\$",
        "/": r"/\allowbreak{}",
        "^": r"\textasciicircum{}", "~": r"\textasciitilde{}",
    }
    return "".join(repl.get(ch, ch) for ch in s)


def latex_header(title):
    return rf"""\documentclass[12pt,a4paper,fontset=none]{{ctexart}}
\usepackage[left=3cm,right=2.5cm,top=2.5cm,bottom=2.5cm]{{geometry}}
\usepackage{{amsmath,amssymb,graphicx,booktabs,longtable,array,enumitem,hyperref,xcolor,float}}
\setmainfont[Path=/usr/share/fonts/dejavu/,BoldFont=DejaVuSans-Bold.ttf,ItalicFont=DejaVuSans-Oblique.ttf]{{DejaVuSans.ttf}}
\setsansfont[Path=/usr/share/fonts/dejavu/,BoldFont=DejaVuSans-Bold.ttf,ItalicFont=DejaVuSans-Oblique.ttf]{{DejaVuSans.ttf}}
\setmonofont[Path=/usr/share/fonts/dejavu/]{{DejaVuSansMono.ttf}}
\setCJKmainfont[Path=/usr/share/fonts/google-droid/]{{DroidSansFallback.ttf}}
\setlength{{\parindent}}{{2em}}
\linespread{{1.5}}
\setlength{{\emergencystretch}}{{3em}}
\sloppy
\hypersetup{{hidelinks}}
\title{{\bfseries {title}}}
\author{{研究生：\underline{{\hspace{{4cm}}}}\quad 指导教师：\underline{{\hspace{{4cm}}}}}}
\date{{二〇二六年十月}}
\begin{{document}}
\maketitle
"""


def text_to_latex(text):
    blocks = []
    for para in text.split("\n\n"):
        if not para.strip():
            continue
        safe = latex_escape(para.strip())
        if para[0].isdigit() and ("、" in para[:3] or "." in para[:3]):
            blocks.append(r"\subsection*{" + safe + "}")
        else:
            blocks.append(safe + "\n\n")
    return "\n".join(blocks)


def generate_tex_files():
    proposal = latex_header("研究生学位（毕业）论文开题报告")
    proposal += rf"""
\begin{{center}}
{{\zihao{{-2}}\bfseries {ZH_TITLE}}}\\[0.6em]
{EN_TITLE}\\[1em]
论文选题项目是否涉密：否
\end{{center}}
\section*{{一、摘要}}
{PROPOSAL_ABSTRACT}
\par\noindent\textbf{{关键词：}}联邦学习；参数高效微调；低秩适配；矩阵级调度；资源异构
\section*{{二、立题依据}}
{text_to_latex(SIGNIFICANCE)}
{text_to_latex(STATE_OF_ART)}
\subsection*{{3、主要参考文献及出处}}
\begingroup\small\setlength{{\parskip}}{{1pt}}\setlength{{\baselineskip}}{{12pt}}
"""
    proposal += "\n\n".join(latex_escape(ref) for ref in REFERENCES[:25])
    proposal += rf"""
\endgroup
\section*{{三、研究方案}}
{text_to_latex(RESEARCH_PLAN)}
\begin{{figure}}[htbp]\centering
\includegraphics[width=.94\textwidth]{{../05_figures/fig_structure/preview/fig2_text_aligned_final.pdf}}
\caption{{FedMATRIX-LoRA总体技术路线。}}\end{{figure}}
{text_to_latex(METHOD_ROUTE)}
下列公式中，$i$、$k$、$j$和$L$分别表示客户端、通信轮次、候选矩阵位置和候选矩阵总数；$m_{{i,j}}^k=1$表示矩阵被激活。$g_{{i,j}}^k$与$v_{{i,j}}^k$分别表示原始FIM近似分数和归一化矩阵效用，$V_i^k$表示整个配置的画像效用。训练显存中的固定项、每矩阵训练状态和位置相关激活分别由$M_i^{{\mathrm{{fixed}}}}$、$q_i(r)c(\mathbf m)$和求和项表示。
\begin{{align*}}
\mathbf m_i^k &= [m_{{i,1}}^k,\ldots,m_{{i,L}}^k],\quad m_{{i,j}}^k\in\{{0,1\}},\\
v_{{i,j}}^k &= g_{{i,j}}^k / \max_{{h,\ell}}g_{{h,\ell}}^k,\\
V_i^k(\mathbf m_i^k) &= \sum_{{j=1}}^L m_{{i,j}}^k v_{{i,j}}^k,\\
M_i(\mathbf m) &= M_i^{{\mathrm{{fixed}}}}+q_i(r)c(\mathbf m)+\sum_{{j=j_{{\min}}}}^L\mu_{{i,j}}.
\end{{align*}}
{text_to_latex(INNOVATION)}
\section*{{四、研究条件与基础}}
{text_to_latex(FOUNDATION)}
\begin{{figure}}[H]\centering
\includegraphics[width=.86\textwidth]{{../05_figures/final/fig1_main_result_matrix.pdf}}
\caption{{已完成实验的多任务最佳指标矩阵。}}\end{{figure}}
\begin{{figure}}[H]\centering
\includegraphics[width=.96\textwidth]{{../05_figures/final/app_full_budget_curves.pdf}}
\caption{{分类任务中的累计预算--性能曲线。}}\end{{figure}}
\begin{{figure}}[H]\centering
\includegraphics[width=.96\textwidth]{{../05_figures/final/app_full_time_curves.pdf}}
\caption{{分类任务中的累计估计时间--性能曲线。}}\end{{figure}}
\section*{{五、论文工作计划及预期成果}}
\begin{{longtable}}{{p{{3.1cm}}p{{9.2cm}}p{{2cm}}}}
\toprule 起止时间 & 工作内容 & 备注\\ \midrule
"""
    schedule = [
        ("2026.10--2026.11", "完善文献调研、问题建模与开题论证", "完成开题"),
        ("2026.11--2027.01", "统一代码、配置清单与运行元数据，补充关键多种子实验", "阶段实验"),
        ("2027.01--2027.02", "开展真实显存与墙钟时间插桩，校准资源模型", "模型校准"),
        ("2027.02--2027.03", "完成大规模与关键消融实验，整理统计结果", "实验定稿"),
        ("2027.03--2027.04", "撰写和修改学位论文方法、系统与实验章节", "论文初稿"),
        ("2027.04--2027.05", "根据导师意见修改论文，完成图表、参考文献与查重", "论文定稿"),
        ("2027.05--2027.06", "准备送审、答辩材料和代码归档", "答辩准备"),
    ]
    for row in schedule:
        proposal += " & ".join(latex_escape(x) for x in row) + r"\\" + "\n"
    proposal += r"\bottomrule\end{longtable}" + "\n" + text_to_latex(EXPECTED)
    proposal += r"\section*{六、签字栏}\noindent 研究生（签名）：\underline{\hspace{5cm}}\hfill 年\quad 月\quad 日\par\vspace{2em}\noindent 指导教师（签名）：\underline{\hspace{5cm}}\hfill 年\quad 月\quad 日\end{document}"
    (OUT / "研究生学位论文开题报告_FedMATRIX-LoRA.tex").write_text(proposal, encoding="utf-8")

    review = latex_header("异构联邦微调中的LoRA资源分配与矩阵级调度研究综述")
    review += rf"""
\begin{{abstract}}
{REVIEW_ABSTRACT}
\end{{abstract}}
\par\noindent\textbf{{关键词：}}联邦学习；参数高效微调；LoRA；异构资源；矩阵级调度
"""
    for heading, paragraphs in REVIEW_SECTIONS:
        review += "\n\\section*{" + latex_escape(heading) + "}\n"
        review += "\n\n".join(latex_escape(x) for x in paragraphs) + "\n"
    review += "\n\\section*{参考文献}\n\\begin{enumerate}[label={[\\arabic*]},leftmargin=2.5em,itemsep=0.2em]\n"
    for ref in REFERENCES:
        ref_body = ref.split("]", 1)[1].strip()
        review += "\\item " + latex_escape(ref_body) + "\n"
    review += "\\end{enumerate}\n\\medskip\\noindent 研究生签名：\\underline{\\hspace{4cm}}\\quad 指导教师签名：\\underline{\\hspace{4cm}}\\quad 提交日期：\\underline{\\hspace{3cm}}\n\\end{document}\n"
    (OUT / "研究生学位论文文献综述_异构联邦LoRA.tex").write_text(review, encoding="utf-8")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    p = generate_proposal()
    r = generate_review()
    generate_tex_files()
    print(p)
    print(r)


if __name__ == "__main__":
    main()
